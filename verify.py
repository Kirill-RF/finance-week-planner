# -*- coding: utf-8 -*-
import datetime as dt, logging
logging.disable(logging.CRITICAL)
import schedula
import formulas as F

FN="План на фин неделю.xlsx"
sol=F.ExcelModel().loads(FN).finish().calculate()
print("calc done")

def cell(sheet,coord):
    key=f"'[{FN.upper()}]{sheet}'!{coord}".upper()
    v=sol.get(key,None)
    return v

import re
def getv(sheet,coord):
    for k,v in sol.items():
        if k.endswith(f"]{sheet}'!{coord}".upper()) and f"[{FN.upper()}]" in k:
            try: return v.value[0,0]
            except Exception: return v
    return None

YEAR=2026
jan1=dt.date(YEAR,1,1)
start=jan1-dt.timedelta(days=(jan1.weekday()-3)%7)
weeks=[]; s=start
while s<=dt.date(YEAR,12,31):
    e=s+dt.timedelta(days=6); weeks.append((s,e)); s=e+dt.timedelta(days=1)
ND=len(weeks)*7
wd_in_month={}
for m in range(1,13):
    wd_in_month[m]=sum(1 for i in range(ND) if (start+dt.timedelta(days=i)).month==m and (start+dt.timedelta(days=i)).weekday()<5)
plans=[5_000_000,4_800_000,5_200_000,5_500_000,5_300_000,6_000_000,6_200_000,6_000_000,5_800_000,6_500_000,7_000_000,8_000_000]

def expected_week(k):  # k 1-based
    a,b=weeks[k-1]
    tot=0.0
    for i in range(7):
        d=a+dt.timedelta(days=i)
        if d.weekday()<5:
            tot+=plans[d.month-1]/wd_in_month[d.month]
    return tot

WK_FIRST=20  # HR=19 -> first data row 20
DAY_FIRST=5
H1rows=ND
DAY_LAST=DAY_FIRST+12*ND-1
H2=DAY_LAST+3
WK_TBL=H2+1

fails=[]
# 1) рабочие дни месяца
for m in range(1,13):
    got=getv("Рабочие дни",f"B{4+m}")
    if abs((got or 0)-wd_in_month[m])>1e-9: fails.append(("wd month",m,got,wd_in_month[m]))
# 2) рабочих дней в неделе
for k in range(1,len(weeks)+1):
    got=getv("Рабочие дни",f"E{WK_FIRST+k-1}")
    exp=sum(1 for i in range(7) if (weeks[k-1][0]+dt.timedelta(days=i)).weekday()<5)
    if abs((got or 0)-exp)>1e-9: fails.append(("wd week",k,got,exp))
# 3) Иванов план недели (блок2 D, строки WK_TBL..WK_TBL+52)
for k in range(1,len(weeks)+1):
    got=getv("Расчёт",f"D{WK_TBL+k-1}")
    exp=expected_week(k)
    if got is None or abs(got-exp)>1.0: fails.append(("week plan",k,got,round(exp,2)))
# 4) стык месяцев: для каждой недели E+F=D и проверка OK
for k in range(1,len(weeks)+1):
    r=WK_TBL+k-1
    d=getv("Расчёт","D%d"%r); e=getv("Расчёт","E%d"%r); f=getv("Расчёт","F%d"%r); chk=getv("Расчёт","G%d"%r)
    if d is None or abs((e+f)-d)>1.0 or chk!="OK": fails.append(("joint",k,d,e,f,chk))
# 5) лист "План на фин. неделю": таблица1 = блок2 (Иванов), активная неделя
act=1+int((dt.date(2026,10,22)-start).days/7)
for k in range(1,len(weeks)+1):
    got=getv("План на фин. неделю",f"{chr(65+k)}10") if k<=25 else None
# better use column letters properly
from openpyxl.utils import get_column_letter
for k in range(1,len(weeks)+1):
    L=get_column_letter(1+k)
    got=getv("План на фин. неделю",f"{L}10")
    exp=expected_week(k)
    if got is None or abs(got-exp)>1.0: fails.append(("wsW t1",k,got,round(exp,2)))
n_act=getv("План на фин. неделю",f"{get_column_letter(len(weeks)+3)}8")
if n_act!=act: fails.append(("active num",n_act,act))
p_act=getv("План на фин. неделю",f"{get_column_letter(len(weeks)+4)}11")  # Иванов строка 11
exp=expected_week(act)
if p_act is None or abs(p_act-exp)>1.0: fails.append(("active plan",p_act,round(exp,2)))
# ИТОГО по году у Иванова = сумма месячных планов
tot_year=sum(getv("План на фин. неделю",f"{get_column_letter(1+k)}10") for k in range(1,len(weeks)+1))
if abs(tot_year-sum(plans))>50: fails.append(("year sum",tot_year,sum(plans)))

print("expected active week:",act,"period:",weeks[act-1])
print("FAILS:",len(fails))
for f in fails[:20]: print(f)
