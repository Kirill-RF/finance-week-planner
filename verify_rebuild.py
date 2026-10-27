# -*- coding: utf-8 -*-
"""Независимая сверка значений в xlsx с бизнес-логикой, посчитанной с нуля."""
import openpyxl, warnings, datetime as dt, calendar as _cal
warnings.filterwarnings('ignore')
F = 'План на фин неделю (обновлённый).xlsx'
wb = openpyxl.load_workbook(F, data_only=True)

ws = wb['Расчёт']
for r in range(1, ws.max_row + 1):
    if ws.cell(r, 2).value == '№ недели':
        h2 = r; break
ok = bad = 0
for rr in range(h2 + 1, h2 + 1 + 71 * 23):
    v = ws.cell(rr, 7).value
    ok += v == 'OK'; bad += v == 'ОШИБКА'
print(f"Контроль стыков (столбец «Проверка A+B=итого»): OK={ok}, ОШИБКА={bad}")

pm = wb['Планы менеджеров']
months = [(2026, m) for m in range(9, 13)] + [(2027, m) for m in range(1, 13)]
# ключи месяцев читаем из шапки листа (строка 3), а не вычисляем — исключает рассинхрон
hdr_keys = [int(pm.cell(3, c).value) for c in range(2, 18)]
plans = {}
order = []
for r in range(6, 29):
    nm = pm.cell(r, 1).value
    order.append(nm)
    vals = {}
    for ci, k in enumerate(hdr_keys):
        v = pm.cell(r, 2 + ci).value
        vals[k] = float(v) if isinstance(v, (int, float)) else 0.0
    plans[nm] = vals

wdays = {}
rd = wb['Рабочие дни']
for r in range(5, 23):
    key = rd.cell(r, 6).value
    if isinstance(key, int):
        y, m = divmod(key, 100)
        wdays[key] = sum(1 for d in range(1, _cal.monthrange(y, m)[1] + 1)
                         if dt.date(y, m, d).weekday() < 5)

first = dt.date(2026, 9, 1)
s0 = first - dt.timedelta(days=(first.weekday() - 3) % 7)
ly, lm = months[-1]
e0 = dt.date(ly, lm, _cal.monthrange(ly, lm)[1])
e0 += dt.timedelta(days=(2 - e0.weekday()) % 7)
weeks = []; s = s0
while s <= e0:
    weeks.append((s, s + dt.timedelta(days=6))); s += dt.timedelta(days=7)

def dayplan(nm, d):
    if d.weekday() >= 5: return 0.0
    k = d.year * 100 + d.month          # ключ ГГГГММ — как в модели файла
    wd = wdays.get(k, 0)
    return plans[nm].get(k, 0.0) / wd if wd else 0.0

err = 0
wf = wb['План на фин. неделю']
mgr_rows = {wf.cell(r, 1).value: r for r in range(10, 33)}
for wi, (a, b) in enumerate(weeks):
    exp_total = 0.0
    for nm in order:
        exp = sum(dayplan(nm, a + dt.timedelta(days=i)) for i in range(7))
        got = wf.cell(mgr_rows[nm], 2 + wi).value or 0
        exp_total += exp
        if abs(got - exp) > 0.02:
            err += 1; print('MISMATCH', nm[:40], 'нед', wi + 1, got, round(exp, 2))
    it = wf.cell(33, 2 + wi).value or 0
    if abs(it - exp_total) > 0.05:
        err += 1; print('ИТОГО mismatch нед', wi + 1)
print(f"Независимая сверка ТАБЛИЦЫ 1: расхождений = {err} (недель {len(weeks)}, менеджеров {len(order)})")

ad = pm.cell(5, 20).value
awn = pm.cell(8, 19).value
idx = [i + 1 for i, (a, b) in enumerate(weeks) if a <= ad.date() <= b]
print(f"Активная дата {ad:%d.%m.%Y}: № недели в файле = {awn}, по календарю = {idx}")
val = pm.cell(11, 20).value
expv = sum(dayplan(order[0], idx[0 - 1][0] if False else weeks[awn - 1][0] + dt.timedelta(days=i)) for i in range(7))
print(f"ТАБЛИЦА 2, первый менеджер: {val:,.2f} vs ожидание {expv:,.2f} -> {'OK' if abs(val - expv) < 0.02 else 'РАСХОЖДЕНИЕ'}")
tot = pm.cell(34, 20).value
expt = sum(sum(dayplan(nm, weeks[awn - 1][0] + dt.timedelta(days=i)) for i in range(7)) for nm in order)
print(f"ТАБЛИЦА 2, ИТОГО: {tot:,.2f} vs ожидание {expt:,.2f} -> {'OK' if abs(tot - expt) < 0.05 else 'РАСХОЖДЕНИЕ'}")
assert err == 0 and bad == 0 and awn in idx, "ЕСТЬ РАСХОЖДЕНИЯ"
print("\nВСЕ ПРОВЕРКИ ПРОЙДЕНЫ")
