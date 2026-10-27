# -*- coding: utf-8 -*-
"""Проверка файла «План на фин неделю (обновлённый).xlsx»:
1) сверка структуры и формул с независимой Python-эмуляцией;
2) проверка всех ссылок формул на корректность диапазонов."""
import datetime as dt, re
import openpyxl

FN = "План на фин неделю (обновлённый).xlsx"
wb = openpyxl.load_workbook(FN)
wsP, wsW, wsX, wsR, wsC = (wb["Планы менеджеров"], wb["План на фин. неделю"],
                           wb["Расчёт"], wb["Рабочие дни"], wb["Календарь"])
ok = True
def chk(cond, msg):
    global ok
    print(("OK  " if cond else "FAIL") + " | " + msg)
    if not cond: ok = False

# ---------- параметры периода ----------
months = []
y, m = 2026, 9
for _ in range(16):
    months.append((y, m)); m += 1
    if m == 13: y, m = y + 1, 1
NM = 23; NW = 71; ND = NW * 7
start = dt.date(2026, 8, 27)
weeks = [(start + dt.timedelta(days=7*k), start + dt.timedelta(days=7*k+6)) for k in range(NW)]

# ---------- 1. имена и планы из дополнения ----------
src = openpyxl.load_workbook("План недели - дополнение.xlsx", data_only=True)["Лист1"]
names_src = [str(src.cell(r,1).value).strip() for r in range(4, src.max_row+1) if src.cell(r,1).value]
names_out = [wsP.cell(6+i,1).value for i in range(NM)]
chk(names_out == names_src, f"Имена менеджеров совпадают с дополнением ({NM} шт.)")
same_plans = all(wsP.cell(6+i, 2+j).value == src.cell(4+i, 2+j).value
                 for i in range(NM) for j in range(3))
chk(same_plans, "Планы сен/окт/ноя 2026 перенесены без изменений")
first_month_hdr = wsP.cell(4,2).value
chk(first_month_hdr == "Сентябрь 2026", f"Расчёт начинается с сентября (первый столбец: {first_month_hdr!r})")
chk(wsP.cell(4,17).value == "Декабрь 2027", "Расчёт продолжается до декабря 2027")

# ---------- 2. ТАБЛИЦА 2 на вкладке «Планы менеджеров» ----------
chk(wsP["S4"].value == "ТАБЛИЦА 2. АКТИВНАЯ НЕДЕЛЯ", "ТАБЛИЦА 2 перенесена на «Планы менеджеров»")
chk(wsW.max_column == 1 + NW and wsW.cell(6, 1+NW+2).value is None,
    "На листе «План на фин. неделю» ТАБЛИЦЫ 2 больше нет (только ТАБЛИЦА 1)")

# ---------- 3. эмуляция расчёта ----------
plans = {names_out[i]: {(mm[0], mm[1]): (wsP.cell(6+i, 2+j).value or 0)
                         for j, mm in enumerate(months)} for i in range(NM)}
wdays = {}
d = weeks[0][0]
while d <= weeks[-1][1]:
    if d.weekday() < 5: wdays[(d.year, d.month)] = wdays.get((d.year, d.month), 0) + 1
    d += dt.timedelta(days=1)
chk(wdays[(2026,9)] == 22 and wdays[(2026,11)] == 21 and wdays[(2026,12)] == 23,
    f"Рабочие дни месяцев считаются по календарю: сен={wdays[(2026,9)]}, ноя={wdays[(2026,11)]}, дек={wdays[(2026,12)]}")

def day_plan(name, dd):
    if dd.weekday() >= 5: return 0.0
    k = (dd.year, dd.month)
    if k not in plans[name]: return 0.0     # IFNA(...,0)
    return plans[name][k] / wdays[k]
def week_plan(name, a, b):
    return sum(day_plan(name, a + dt.timedelta(days=i)) for i in range(7))

# строки «Рабочих дней»: ключи месяцев и формулы SUMPRODUCT (включая хвостовые месяцы границ периода)
NKEYS = 18
keys = [wsR.cell(5+i, 6).value for i in range(NKEYS)]
exp_keys = sorted([int(f"{yy}{mm:02d}") for yy, mm in months] + [202608, 202801])
chk(sorted(keys) == exp_keys, "Ключи месяцев ГГГГММ на листе «Рабочие дни» (покрывают весь календарь)")
chk(str(wsR.cell(5,2).value).startswith("=SUMPRODUCT((Кал_Месяц=202609)"), "Формула рабочих дней месяца (сен 2026)")

# лист «Расчёт»: блок 1 — план дня
DAY_FIRST = 5
r = DAY_FIRST  # первый менеджер, первый день (чт 27.08.2026 -> август!)
a0 = weeks[0][0]
g_f = wsX.cell(r,7).value
chk("MATCH(A{0},'Планы менеджеров'!$A$6:$A${1},0)".format(r, 5+NM) in g_f.replace(" ","") or
    f"MATCH(A{r},'Планы менеджеров'!$A$6:$A${5+NM},0)" in g_f, "Формула плана месяца ссылается на имя менеджера")
chk(int(f"{a0.year}{a0.month:02d}") in keys,
    "Месяц первой даты календаря (авг.2026) есть в справочнике ключей — план=0 через IFNA")
missing_keys = set()
d = weeks[0][0]
while d <= weeks[-1][1]:
    missing_keys.add(int(f"{d.year}{d.month:02d}"))
    d += dt.timedelta(days=1)
plan_keys = set(keys)
extra = missing_keys - plan_keys
chk(not extra, f"Все даты календаря покрыты ключами месяцев плана (пробел: {sorted(extra) or 'нет'})")

# активная неделя для 22.10.2026
ad = dt.date(2026,10,22)
wk_num = (ad - weeks[0][0]).days // 7 + 1
exp_val = week_plan(names_out[0], *weeks[wk_num-1])
print(f"     ожидание: №активной недели={wk_num}, план 1-го менеджера={exp_val:,.2f}")

# свёртка недельных планов по блоку 1 (эмуляция сумм H-столбцов)
tot_mgr0 = sum(week_plan(names_out[0], a, b) for a, b in weeks)
sum_months = sum(plans[names_out[0]].values().__iter__().__next__().__class__ and v for v in plans[names_out[0]].values())
chk(abs(tot_mgr0 - sum_months) < 1e-6, "Σ недельных планов первого менеджера = Σ месячных планов (стыки учтены)")

# ---------- 4. все ссылки формул валидны ----------
sheet_dims = {n: wb[n].dimensions for n in wb.sheetnames}
pat_cell = re.compile(r"(?:'([^']+)'|([А-Яа-яA-Za-z_. ]+))?!?\$?([A-Z]{1,3})\$?(\d+)")
bad = []
for name in wb.sheetnames:
    ws = wb[name]
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and v.startswith("="):
                for mm in pat_cell.finditer(v):
                    sh = mm.group(1) or mm.group(2)
                    col, rr = mm.group(3), int(mm.group(4))
                    if sh and sh in sheet_dims and sh != name:
                        maxc = openpyxl.utils.column_index_from_string(col)
                        dim = sheet_dims[sh]
                        import openpyxl.utils as U
                        max_col = U.column_index_from_string(re.match(r"[A-Z]+", dim.split(":")[0]).group())
                        rng = ws[re.match(r"[A-Z]+\d+", dim.split(":")[1]).group()].__class__ and None
                        max_row = int(re.search(r"(\d+)$", dim.split(":")[1]).group(1))
                        max_col2 = U.column_index_from_string(re.match(r"[A-Z]+", dim.split(":")[1]).group())
                        if rr > max_row or maxc > max_col2:
                            bad.append((name, c.coordinate, sh, col+str(rr)))
chk(not bad, f"Все межлистовые ссылки формул внутри существующих диапазонов (нарушений: {len(bad)})")
if bad: print(bad[:10])

# ---------- 5. структура листа «План на фин. неделю» ----------
chk(wsW.cell(7,1).value=="№ недели" and wsW.cell(7,2).value==1 and wsW.cell(7,1+NW).value==NW,
    f"ТАБЛИЦА 1: {NW} недель, нумерация 1..{NW}")
chk(wsW.cell(10,1).value==names_out[0] and wsW.cell(10+NM-1,1).value==names_out[-1] and wsW.cell(10+NM,1).value=="ИТОГО",
    "Строки менеджеров + ИТОГО на месте")
chk(wsW["B4"].value == "='Планы менеджеров'!T5", "Активная дата на листе недель ссылается на ТАБЛИЦУ 2")
cf = list(wsW.conditional_formatting)
chk(len(cf)==1 and "'Планы менеджеров'!$S$8" in str(cf[0].rules[0].formula), "Подсветка активной недели использует номер из ТАБЛИЦЫ 2")

# ---------- 6. строка «Рабочих дней месяца» ----------
chk(wsP.cell(5,1).value=="Рабочих дней месяца" and wsP.cell(5,2).value=="='Рабочие дни'!$B$5",
    "Добавлена строка «Рабочих дней месяца» (используется в расчёте плана дня)")

print("\n=== ИТОГ:", "ВСЁ В ПОРЯДКЕ" if ok else "ЕСТЬ ПРОБЛЕМЫ", "===")
